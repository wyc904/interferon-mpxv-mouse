# -*- coding: utf-8 -*-
"""读取 3-1009 批次的样本信息（只读，不改动源文件）。

内容：分组信息.docx / 肺与脾组织样本编号对照表.xlsx / MPXV_RB_counts_gene(1).txt 表头
输出：2-script/_source_meta_dump.txt (UTF-8)
"""
from pathlib import Path

SRC = Path(r"D:\Desktop\临时保存\5-干扰素小鼠\3-1009\1-源")
OUT = SRC.parent / "2-script" / "_source_meta_dump.txt"

out = []


def w(s=""):
    out.append(str(s))


w("3-1009 源数据信息汇总 (inspect_source_meta.py 生成)")
w()

# ---------- 1. 分组信息.docx ----------
w("=" * 78)
w("## 1. 分组信息.docx")
w("=" * 78)
try:
    import docx

    doc = docx.Document(SRC / "分组信息.docx")
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if t:
            try:
                style = p.style.name
            except Exception:
                style = "?"
            w(f"[P{i:03d}|{style}] {t}")
    for ti, tbl in enumerate(doc.tables):
        w()
        w(f"--- docx 表格{ti + 1}: {len(tbl.rows)}行 x {len(tbl.columns)}列 ---")
        for ri, r in enumerate(tbl.rows):
            cells = [c.text.strip().replace("\n", " ") for c in r.cells]
            w(f"R{ri:02d}: " + " | ".join(cells))
except Exception as e:
    w(f"!! docx 读取失败: {type(e).__name__}: {e}")

# ---------- 2. 肺与脾组织样本编号对照表.xlsx ----------
w()
w("=" * 78)
w("## 2. 肺与脾组织样本编号对照表.xlsx")
w("=" * 78)
try:
    import openpyxl

    wb = openpyxl.load_workbook(SRC / "肺与脾组织样本编号对照表.xlsx", data_only=True)
    w(f"sheets = {wb.sheetnames}")
    for ws in wb.worksheets:
        w()
        w(f"--- sheet [{ws.title}]  {ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column} ---")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            if any(cells):
                w(" | ".join(cells))
except Exception as e:
    w(f"!! xlsx 读取失败: {type(e).__name__}: {e}")

# ---------- 3. MPXV_RB_counts_gene(1).txt ----------
w()
w("=" * 78)
w("## 3. MPXV_RB_counts_gene(1).txt")
w("=" * 78)
txt = SRC / "MPXV_RB_counts_gene(1).txt"
comment_lines, header, n_data = [], None, 0
gene_names, first_rows = [], []
with open(txt, "r", encoding="utf-8", errors="replace") as f:
    for ln in f:
        if ln.startswith("#"):
            comment_lines.append(ln.rstrip("\n"))
            continue
        if header is None:
            header = ln.rstrip("\n")
            continue
        n_data += 1
        if len(first_rows) < 2:
            first_rows.append(ln.rstrip("\n"))
        gene_names.append(ln.split("\t", 1)[0])

for c in comment_lines:
    w("[注释] " + c)
w(f"数据行数(基因行) = {n_data};  唯一基因名 = {len(set(gene_names))}")
cols = header.split("\t")
w(f"表头列数 = {len(cols)} (首列 = {cols[0]!r})")
w(f"样本列共 {len(cols) - 1} 个:")
for i, c in enumerate(cols[1:], 1):
    w(f"  {i:02d}. {c}")
w()
w("前 2 行数据(每行截断 150 字符):")
for d in first_rows:
    w("  " + d[:150])

OUT.write_text("\n".join(out), encoding="utf-8")
print("WROTE:", str(OUT))
print("total_lines:", len(out))
