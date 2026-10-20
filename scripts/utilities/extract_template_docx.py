# -*- coding: utf-8 -*-
"""抽取 分析模板.docx 的正文与嵌入图片（只读模板，不改动）。"""
import zipfile
from pathlib import Path

import docx

TPL = Path(r"D:\Desktop\临时保存\5-干扰素小鼠\3-1009\3-分析模板")
OUT = Path(r"D:\Desktop\临时保存\5-干扰素小鼠\3-1009\2-script\_template_extract")
OUT.mkdir(exist_ok=True)
MEDIA = OUT / "media"
MEDIA.mkdir(exist_ok=True)

# 1. 正文与表格
doc = docx.Document(TPL / "分析模板.docx")
lines = []
for i, p in enumerate(doc.paragraphs):
    t = (p.text or "").strip()
    # 记录标题级别/样式，非空即写
    if t:
        lines.append(f"[P{i:03d}|{p.style.name}] {t}")
for ti, tbl in enumerate(doc.tables):
    lines.append(f"\n--- 表格{ti+1}: {len(tbl.rows)}x{len(tbl.columns)} ---")
    for r in tbl.rows:
        lines.append(" | ".join(c.text.strip().replace("\n", " ") for c in r.cells))
(OUT / "docx_text.txt").write_text("\n".join(lines), encoding="utf-8")

# 2. 提取嵌入图片
with zipfile.ZipFile(TPL / "分析模板.docx") as z:
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    for n in media:
        data = z.read(n)
        (MEDIA / Path(n).name).write_bytes(data)

print("paras/tables lines:", len(lines))
print("media files:", len(media))
for n in media:
    print("  ", n)
