# -*- coding: utf-8 -*-
"""解构热图：按 DEG_Source 分块输出每基因各组平均 Z 值，供解读。"""
import statistics
from collections import OrderedDict
import openpyxl

XLSX = r"D:\Desktop\临时保存\5-干扰素小鼠\3-1009\4-结果\BALBc-肺\00-表格\热图数据_BALBc_Lung.xlsx"
OUT = r"D:\Desktop\临时保存\5-干扰素小鼠\3-1009\2-script\_cache\_heat_BALBc_Lung_summary.txt"

wb = openpyxl.load_workbook(XLSX, read_only=True)

# DEGinfo: gene -> (source, log2FC, padj)
deg = OrderedDict()
ws = wb["DEGinfo"]
ws.reset_dimensions()
rows = ws.iter_rows(values_only=True)
hdr = list(next(rows))
gi, ci = hdr.index("Geneid"), hdr.index("comparison")
li, pi = hdr.index("log2FoldChange"), hdr.index("padj")
for r in rows:
    deg[r[gi]] = (r[ci], r[li], r[pi])

# Zscore: gene -> {sample: z}
zmap = {}
ws = wb["Zscore"]
ws.reset_dimensions()
rows = ws.iter_rows(values_only=True)
hdr = list(next(rows))
cols = hdr[1:]
for r in rows:
    zmap[r[0]] = dict(zip(cols, r[1:]))
wb.close()


def mz(samples, gr):
    vals = [v for c, v in samples.items() if c.startswith(gr + "-")]
    return statistics.mean(vals) if vals else float("nan")


lines = []
lines.append("gene\tsource\tlog2FC\tpadj\tZ_G1\tZ_G2\tZ_G3\tZ_G4")
for g, (comp, lfcv, pv) in deg.items():
    z = zmap[g]
    lines.append(f"{g}\t{comp}\t{lfcv:.2f}\t{pv:.2g}\t{mz(z,'G1'):.2f}\t{mz(z,'G2'):.2f}\t{mz(z,'G3'):.2f}\t{mz(z,'G4'):.2f}")

# 按来源比较组分块汇总
lines.append("\n=== BLOCKS ===")
blocks = OrderedDict()
for g, (comp, lfcv, pv) in deg.items():
    blocks.setdefault(comp, []).append((g, lfcv, pv, zmap[g]))
for comp, genes in blocks.items():
    means = {gr: statistics.mean([mz(z, gr) for _, _, _, z in genes]) for gr in ("G1", "G2", "G3", "G4")}
    top = sorted(genes, key=lambda x: -abs(x[1]))[:8]
    lines.append(f"\n[{comp}]  n={len(genes)}  group-meanZ  G1={means['G1']:.2f}  G2={means['G2']:.2f}  G3={means['G3']:.2f}  G4={means['G4']:.2f}")
    for g, lfcv, pv, z in top:
        lines.append(f"    {g:<16} log2FC={lfcv:+.2f} padj={pv:.1e}  Z: G1={mz(z,'G1'):+.2f} G2={mz(z,'G2'):+.2f} G3={mz(z,'G3'):+.2f} G4={mz(z,'G4'):+.2f}")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("WROTE", OUT, len(lines), "lines")
